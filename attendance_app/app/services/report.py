"""Excel report of one class session (openpyxl)."""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..db import EventLog
from . import attendance, storage

NAVY, GOLD, CREAM = "1B2A41", "C9A96E", "F7F1E5"
STATUS_FILL = {"present": "E3F1E7", "late": "FFF1D6", "absent": "F9E0E0"}
EVENT_TH = {"spoof": "สงสัยภาพปลอม (ไม่เช็คชื่อ)", "confirm": "อาจารย์ยืนยันตัวตน", "reject": "อาจารย์ปฏิเสธ",
            "manual_add": "อาจารย์เช็คชื่อเอง", "manual_remove": "อาจารย์ยกเลิกการเช็คชื่อ"}
thin = Side(style="thin", color="D9D2C3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _header(ws, row, titles, widths):
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF", name="Tahoma")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = width


def build_excel(db, session):
    section = session.section_obj
    rows = attendance.roster(db, session)
    counts = attendance.summary(rows)
    wb = Workbook()

    # ---------------- sheet 1: attendance
    ws = wb.active
    ws.title = "สรุปการเข้าเรียน"
    ws["A1"] = f"{section.course_code} {section.course_name}".strip()
    ws["A1"].font = Font(bold=True, size=16, color=NAVY, name="Tahoma")
    ws["A2"] = (f"Section {section.section}  ·  เทอม {section.term}  ·  วันที่ {session.scheduled_start:%d/%m/%Y}  ·  "
                f"เริ่มคาบ {session.scheduled_start:%H:%M}  ·  สายหลัง {session.late_minutes} นาที")
    ws["A3"] = (f"ทั้งหมด {counts['total']} คน  ·  มา {counts['present']}  ·  สาย {counts['late']}  ·  "
                f"ขาด {counts['absent']}")
    ws["A3"].font = Font(bold=True, color="8A6D3B", name="Tahoma")
    titles = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ", "สถานะ", "เวลาเช็คชื่อ", "วิธี", "ความมั่นใจ (similarity)", "หมายเหตุ"]
    _header(ws, 5, titles, [7, 16, 28, 10, 14, 20, 22, 30])
    for i, r in enumerate(rows, start=1):
        note = []
        if r["flagged"]:
            note.append("anti-spoof เตือนว่าอาจเป็นภาพปลอม")
        if r["risk"]:
            note.append("กลุ่มเสี่ยง (แฝด/หน้าคล้าย)")
        values = [i, r["code"], r["name"], r["status_th"], r["checked_at"] or "-", r["method_th"] or "-",
                  r["similarity"] if r["similarity"] is not None else "-", ", ".join(note)]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=5 + i, column=col, value=value)
            cell.border = BORDER
            cell.font = Font(name="Tahoma")
            cell.alignment = Alignment(horizontal="left" if col in (3, 8) else "center")
            if col == 4:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL[r["status"]])
    ws.freeze_panes = "A6"

    # ---------------- sheet 2: event log (teacher confirmations, spoof, manual)
    log = wb.create_sheet("บันทึกเหตุการณ์")
    _header(log, 1, ["เวลา", "เหตุการณ์", "นักศึกษา", "รายละเอียด", "ไฟล์ภาพหลักฐาน"], [12, 30, 28, 50, 45])
    events = db.query(EventLog).filter_by(session_id=session.id).order_by(EventLog.at).all()
    for i, e in enumerate(events, start=2):
        detail = json.loads(e.detail or "{}")
        text = ""
        if "candidates" in detail:
            text = "ตัวเลือก: " + ", ".join(f"{c['name']} ({c['similarity']})" for c in detail["candidates"])
        elif "spoof_score" in detail:
            text = f"P(spoof) {detail['spoof_score']}, similarity {detail.get('similarity')}"
            if detail.get("reason_th"):
                text = f"{detail['reason_th']} | {text}, device {detail.get('device_score')}"
        if e.reviewed:
            text += f"  [อาจารย์ตรวจแล้ว: {'คนจริง' if e.reviewed == 'live' else 'ภาพปลอมจริง'}]"
        values = [e.at.strftime("%H:%M:%S"), EVENT_TH.get(e.kind, e.kind), e.student.name if e.student else "-",
                  text, str(storage.MEDIA / e.image_path) if e.image_path else ""]
        for col, value in enumerate(values, start=1):
            cell = log.cell(row=i, column=col, value=value)
            cell.border = BORDER
            cell.font = Font(name="Tahoma")

    ws.sheet_properties.tabColor = GOLD
    log.sheet_properties.tabColor = NAVY
    name = storage.safe_name(f"{section.course_code or section.course_name}_sec{section.section}")
    out = storage.EXPORTS / f"{name}_{session.scheduled_start:%Y%m%d_%H%M}_s{session.id}.xlsx"
    wb.save(out)
    return out
