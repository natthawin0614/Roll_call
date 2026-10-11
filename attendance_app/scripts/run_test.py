"""Test the system on your own photos AND videos, and compare with an answer key (Excel).

    python scripts\\run_test.py                         <- uses ..\\image_pyiot\\test set
    python scripts\\run_test.py "D:\\my clips"           <- any folder with photos / videos + one .xlsx

Answer key (.xlsx in the same folder), first sheet:
    | image_No | Member1 | Member2 | ... | Expect (optional) |
    | image1   | Bus     | Tan     |     |                   |   -> 01.jpg / 1.jpg / image1.jpg
    | clip2    | Ploy    |         |     | spoof             |   -> clip2.mp4: a phone attack, nobody may be checked in
  - the first column is matched to the file name (same name, or the same number: image1 = 01.jpg)
  - "A/B" = either one is fine (twins)
  - Expect = live (default) or spoof
Nicknames -> enrolled names: names.csv in the same folder (created the first time; edit it for new people).
People in the answer key that are not enrolled are listed but not counted as "missed".

Photos are judged like one camera frame (no head-turn possible).
Videos are played through the SAME logic as the live check-in page (tracking, 3-layer anti-spoof,
head-turn challenge, N frames in a row), using the video's own clock.
Output: ..\\test_results\\<time>\\  annotated photos / videos + report.xlsx
"""
import argparse
import csv
import re
import sys
import time
import types
from collections import Counter
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from app.config import config  # noqa: E402
from app.services import live  # noqa: E402
from app.services.enroll import scan_folder  # noqa: E402
from app.services.matcher import Gallery, Person  # noqa: E402
from app.vision.pipeline import get_pipeline, live_quality_issue, read_image  # noqa: E402
from eval_testset import EN2TH  # noqa: E402

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
VIDEO_EXT = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"}
SKIP_ENROL = {"พิชหรือวุฒิแยกโครตยาก"}
FONT_FILE = Path(r"C:\Windows\Fonts\LeelaUIb.ttf")
COLORS = {"success": (87, 192, 143), "already": (201, 169, 110), "checking": (122, 167, 220), "ambiguous": (229, 173, 72),
          "unknown": (152, 163, 182), "spoof": (228, 107, 107), "adjust": (152, 163, 182), "challenge": (229, 173, 72),
          "live": (87, 192, 143)}
STATUS_TH = {"success": "เช็คชื่อ", "already": "เช็คแล้ว", "checking": "กำลังยืนยัน", "ambiguous": "ไม่แน่ใจ",
             "unknown": "ไม่รู้จัก", "spoof": "ปลอม", "adjust": "มองตรง", "challenge": "หันหน้า", "live": "คนจริง"}


def font(size):
    return ImageFont.truetype(str(FONT_FILE), size) if FONT_FILE.exists() else ImageFont.load_default()


# ---------------------------------------------------------------- inputs
def build_gallery(folder, pipeline):
    _, groups = scan_folder(folder)
    people = []
    for (code, name), paths in groups.items():
        if name in SKIP_ENROL:
            continue
        vectors = [f[0].embedding for f in (pipeline.analyze(read_image(p.read_bytes())) for p in paths) if f]
        if vectors:
            people.append(Person(len(people), name, code, False, np.array(vectors, np.float32), None))
    for i, a in enumerate(people):                       # automatic risk group, as in the app
        for b in people[i + 1:]:
            if float((a.vectors @ b.vectors.T).max()) >= config["enroll"]["similar_warn"]:
                a.risk = b.risk = True
    return Gallery(people)


def load_names(folder):
    path = folder / "names.csv"
    if not path.exists():
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["nickname", "enrolled_name"])
            w.writerows(EN2TH.items())
            w.writerows([("Pitch", "พิช"), ("Wuth", "วุต")])
        print(f"created {path} - add a row for every new nickname used in the answer key")
    with open(path, encoding="utf-8-sig") as f:
        return {r["nickname"].strip().lower(): r["enrolled_name"].strip() for r in csv.DictReader(f) if r["nickname"]}


def number_of(text):
    m = re.findall(r"\d+", str(text))
    return int(m[-1]) if m else None


def options_of(nickname, names):
    """'Pitch/Wuth' -> ['พิช', 'วุต'] (either one counts)."""
    whole = names.get(nickname.strip().lower())
    if whole:
        return [x.strip() for x in whole.split("/")]
    return [names.get(x.strip().lower(), x.strip()) for x in nickname.split("/")]


def load_answers(folder, names):
    files = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        print("no .xlsx answer key -> results only, no comparison")
        return {}
    ws = openpyxl.load_workbook(files[0], data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip().lower() for h in rows[0]]
    expect_col = next((i for i, h in enumerate(header) if h in ("expect", "type", "label", "ประเภท")), None)
    answers = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        people = [str(v).strip() for i, v in enumerate(row[1:], start=1) if v and i != expect_col]
        expect = str(row[expect_col]).strip().lower() if expect_col is not None and row[expect_col] else "live"
        answers[str(row[0]).strip()] = {
            "expect": "spoof" if expect in ("spoof", "fake", "ปลอม") else "live",
            "people": [{"key": p, "options": options_of(p, names)} for p in people]}
    print(f"answer key: {files[0].name} ({len(answers)} rows)")
    return answers


def answer_for(path, answers):
    if path.stem in answers:
        return answers[path.stem]
    n = number_of(path.stem)
    for key, value in answers.items():
        if n is not None and number_of(key) == n:
            return value
    return None


def compare(found, answer, enrolled):
    """found: names checked in. Returns per-file result."""
    res = {"correct": [], "wrong": [], "missed": [], "not_enrolled": []}
    if answer is None:
        return res
    if answer["expect"] == "spoof":
        res["wrong"] = list(found)                       # anyone checked in from an attack video = wrong
        return res
    left = list(found)
    for person in answer["people"]:
        hit = next((n for n in left if n in person["options"]), None)
        if hit:
            res["correct"].append(hit)
            left.remove(hit)
        elif any(o in enrolled for o in person["options"]):
            res["missed"].append("/".join(person["options"]))
        else:
            res["not_enrolled"].append(person["key"])
    res["wrong"] = left
    return res


# ---------------------------------------------------------------- photos
def draw_labels(image_bgr, items, scale=1.0, size=24):
    """items: (box, color, label)"""
    canvas = Image.fromarray(cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB))
    d = ImageDraw.Draw(canvas)
    f = font(size)
    for box, color, label in items:
        x1, y1, x2, y2 = [v * scale for v in box]
        d.rectangle([x1, y1, x2, y2], outline=color, width=max(2, size // 6))
        if label:
            tb = d.textbbox((x1, max(0, y1 - size - 8)), label, font=f)
            d.rectangle([tb[0] - 3, tb[1] - 2, tb[2] + 3, tb[3] + 2], fill=color)
            d.text((x1, max(0, y1 - size - 8)), label, font=f, fill=(255, 255, 255))
    return cv2.cvtColor(np.array(canvas), cv2.COLOR_RGB2BGR)


def zone_of(face):
    s = config["antispoof"]
    if s["device_check"] and face.device_score >= s["device_conf"]:
        return "spoof"
    if face.spoof_score <= s["live_below"]:
        return "live"
    return "spoof" if face.spoof_score >= s["spoof_above"] else "challenge"


def run_photo(path, pipeline, gallery, out_dir):
    image = read_image(path.read_bytes())
    faces = pipeline.analyze(image, min_face_px=40, devices=True)
    found, items, notes = [], [], Counter()
    for fc in faces:
        d = gallery.identify(fc.embedding)
        name = gallery.by_id[d.student_id].name if d.student_id is not None else None
        zone = zone_of(fc)
        notes[zone] += 1
        if d.status == "match" and zone != "spoof" and name not in found:
            found.append(name)
        label = (name if d.status == "match" else f"?{name}" if d.status == "ambiguous" else "ไม่รู้จัก") + \
            f" {d.similarity:.2f} | {STATUS_TH[zone]}"
        color = COLORS["unknown"] if d.status != "match" else COLORS[zone]
        items.append((fc.box, color, label))
        if fc.device_box and fc.device_score >= config["antispoof"]["device_conf"]:
            items.append((fc.device_box, COLORS["spoof"], f"มือถือ/จอ {fc.device_score:.2f}"))
    scale = 1600 / max(image.shape[:2])
    small = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    cv2.imencode(".jpg", draw_labels(small, items, scale))[1].tofile(str(out_dir / f"{path.stem}.jpg"))
    return found, {"faces": len(faces), "zones": dict(notes)}


# ---------------------------------------------------------------- videos (same logic as the live page)
class Clock:
    t = 0.0

    @classmethod
    def time(cls):
        return cls.t


class VideoSession(live.LiveSession):
    """Live check-in without the database: who would have been checked in from this video."""

    def __init__(self, gallery):
        super().__init__(0)
        self.gallery = gallery
        self.checked = {}          # name -> second
        self.spoofs = []           # (second, reason, look-alike name)
        self.challenges = 0

    def step(self, frame):
        a, s = config["attendance"], config["antispoof"]
        mode = s["mode"]
        faces = get_pipeline().analyze(frame, min_face_px=a["min_face_px"], devices=mode != "off" and s["device_check"])
        out = []
        for face, track in zip(faces, self._assign(faces)):
            issue = live_quality_issue(face, frame.shape)
            d = self.gallery.identify(face.embedding)
            if not issue and track.challenge is None:
                track.history.append((d.status, d.student_id))
                track.spoof_scores.append(face.spoof_score)
            had_challenge = track.challenge is not None
            self._update_liveness(track, face, issue, mode)
            if track.challenge is not None and not had_challenge:
                self.challenges += 1
            name = self.gallery.by_id[d.student_id].name if d.student_id is not None else None
            n = a["confirm_frames"]
            recent = list(track.history)[-n:]
            same = len(recent) == n and all(x == "match" for x, _ in recent) and len({i for _, i in recent}) == 1
            if track.liveness == "spoof" and mode == "block":
                status = "spoof"
                last = getattr(track, "evidence_at", None)
                if last is None or Clock.t - last > a["evidence_cooldown_sec"]:
                    track.evidence_at = Clock.t
                    self.spoofs.append((round(Clock.t, 1), live.SPOOF_REASON_TH.get(track.spoof_reason, ""), name))
            elif track.checked_student is not None:
                status, name = "already", self.gallery.by_id[track.checked_student].name
            elif track.challenge is not None:
                c = track.challenge
                status = "challenge"
                name = f"หัน{'ซ้าย' if c['dir'] == 'left' else 'ขวา'} {c['progress']:.0%}"
            elif issue:
                status, name = "adjust", issue
            elif track.liveness is None:
                status = "checking"
            elif same:
                status = "success"
                track.checked_student = d.student_id
                self.checked.setdefault(name, round(Clock.t, 1))
            else:
                status = {"ambiguous": "ambiguous", "unknown": "unknown", "empty": "unknown"}.get(d.status, "checking")
            out.append((face, status, name, d.similarity))
        return out


def run_video(path, gallery, out_dir, fps):
    cap = cv2.VideoCapture(str(path))
    src_fps = cap.get(cv2.CAP_PROP_FPS) or 30
    step = max(1, round(src_fps / fps))
    width = config["camera"]["send_width"]
    session = VideoSession(gallery)
    real_time = live.time
    live.time = types.SimpleNamespace(time=Clock.time)        # the challenge timer follows the video
    writer, i = None, 0
    try:
        while True:
            ok, frame = cap.read()
            if not ok:
                break
            if i % step == 0:
                Clock.t = i / src_fps
                if frame.shape[1] > width:
                    frame = cv2.resize(frame, None, fx=width / frame.shape[1], fy=width / frame.shape[1],
                                       interpolation=cv2.INTER_AREA)
                results = session.step(frame)
                items = [(f.box, COLORS[st], f"{nm or ''} {STATUS_TH[st]}".strip()) for f, st, nm, _ in results]
                items += [(f.device_box, COLORS["spoof"], "มือถือ/จอ") for f, st, _, _ in results
                          if f.device_box and f.device_score >= config["antispoof"]["device_conf"]]
                shown = draw_labels(frame, items, size=22)
                cv2.putText(shown, f"{Clock.t:5.1f}s", (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.9, (255, 255, 255), 2)
                if writer is None:
                    writer = cv2.VideoWriter(str(out_dir / f"{path.stem}.mp4"), cv2.VideoWriter_fourcc(*"mp4v"), fps,
                                             (shown.shape[1], shown.shape[0]))
                writer.write(shown)
            i += 1
    finally:
        live.time = real_time
        cap.release()
        if writer:
            writer.release()
    found = list(session.checked)
    return found, {"seconds": round(i / src_fps, 1), "checked_at": session.checked, "spoofs": session.spoofs,
                   "challenges": session.challenges}


# ---------------------------------------------------------------- report
def write_excel(rows, totals, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ผลทดสอบ"
    head = ["ไฟล์", "ชนิด", "ควรเป็น", "เฉลย (ลงทะเบียนแล้ว)", "ระบบเช็คชื่อ", "ถูก", "ผิดคน", "หาไม่เจอ",
            "ไม่ได้ลงทะเบียน", "รายละเอียด", "ผ่าน?"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A1839")
    for r in rows:
        ws.append([r["file"], r["kind"], r["expect"], r["truth"], ", ".join(r["found"]), len(r["res"]["correct"]),
                   ", ".join(r["res"]["wrong"]), ", ".join(r["res"]["missed"]), ", ".join(r["res"]["not_enrolled"]),
                   r["detail"], "ผ่าน" if r["pass"] else "ไม่ผ่าน" if r["pass"] is not None else "-"])
        fill = "DFF3E8" if r["pass"] else "FBE3E3" if r["pass"] is not None else "FFFFFF"
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    for k, v in totals.items():
        ws.append([k, v])
    for col, w in zip("ABCDEFGHIJK", (14, 8, 9, 40, 40, 6, 22, 22, 22, 45, 9)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", nargs="?", type=Path, default=ROOT.parent / "image_pyiot" / "test set")
    ap.add_argument("--gallery", type=Path, default=ROOT.parent / "image_pyiot", help="enrolment photos (name = file name)")
    ap.add_argument("--out", type=Path, default=ROOT.parent / "test_results")
    ap.add_argument("--fps", type=float, default=6, help="video frames per second to process (the live page does ~5-10)")
    args = ap.parse_args()

    out_dir = args.out / datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir.mkdir(parents=True, exist_ok=True)
    pipeline = get_pipeline()
    t0 = time.time()
    gallery = build_gallery(args.gallery, pipeline)
    enrolled = {p.name for p in gallery.people}
    print(f"enrolled: {len(gallery.people)} people from {args.gallery} | risk group: {[p.name for p in gallery.people if p.risk]}")
    names = load_names(args.folder)
    answers = load_answers(args.folder, names)

    rows = []
    files = sorted(p for p in args.folder.iterdir() if p.suffix.lower() in IMAGE_EXT | VIDEO_EXT)
    for path in files:
        is_video = path.suffix.lower() in VIDEO_EXT
        if is_video:
            found, info = run_video(path, gallery, out_dir, args.fps)
            detail = (f"{info['seconds']}s · เช็คชื่อที่วินาที {info['checked_at'] or '-'} · ให้หันหน้า {info['challenges']} ครั้ง"
                      + (f" · ปลอม: {info['spoofs']}" if info["spoofs"] else ""))
        else:
            found, info = run_photo(path, pipeline, gallery, out_dir)
            detail = f"{info['faces']} หน้า · " + ", ".join(f"{STATUS_TH[k]} {v}" for k, v in info["zones"].items())
        answer = answer_for(path, answers)
        res = compare(found, answer, enrolled)
        if answer is None:
            ok, expect, truth = None, "-", "(ไม่มีในเฉลย)"
        elif answer["expect"] == "spoof":
            ok, expect, truth = not found, "ปลอม", "ห้ามเช็คชื่อใคร"
        else:
            ok, expect = not res["wrong"] and not res["missed"], "คนจริง"
            truth = ", ".join("/".join(p["options"]) for p in answer["people"] if any(o in enrolled for o in p["options"]))
        rows.append({"file": path.name, "kind": "วิดีโอ" if is_video else "รูป", "expect": expect, "truth": truth,
                     "found": found, "res": res, "detail": detail, "pass": ok})
        mark = {True: "PASS", False: "FAIL", None: "----"}[ok]
        print(f"[{mark}] {path.name:12s} found: {', '.join(found) or '-'}"
              + (f" | wrong: {', '.join(res['wrong'])}" if res["wrong"] else "")
              + (f" | missed: {', '.join(res['missed'])}" if res["missed"] else "") + f"   ({detail})")

    judged = [r for r in rows if r["pass"] is not None]
    tp = sum(len(r["res"]["correct"]) for r in rows)
    fp = sum(len(r["res"]["wrong"]) for r in rows)
    fn = sum(len(r["res"]["missed"]) for r in rows)
    prec, rec = tp / max(tp + fp, 1), tp / max(tp + fn, 1)
    totals = {"ไฟล์ที่ผ่าน": f"{sum(r['pass'] for r in judged)}/{len(judged)}",
              "เช็คชื่อถูกคน (TP)": tp, "เช็คชื่อผิดคน (FP)": fp, "หาไม่เจอ (FN)": fn,
              "Precision": round(prec, 3), "Recall": round(rec, 3), "F1": round(2 * prec * rec / max(prec + rec, 1e-9), 3)}
    write_excel(rows, totals, out_dir / "report.xlsx")
    print("\n" + " | ".join(f"{k} {v}" for k, v in totals.items()))
    print(f"results -> {out_dir}  ({time.time() - t0:.0f}s)")


if __name__ == "__main__":
    main()
